#!/usr/bin/env python3
"""
Generates lib/data/fake/seed_data.dart for Learn Ateso from the real
vocabulary spreadsheet filled in by the Ateso-speaking groupmate.

Usage (run from this content/ folder):
  python generate_seed_data.py [vocab.xlsx] [sentence_bank.xlsx] [out.dart]

Defaults:
  vocab.xlsx          Learn_Ateso_Beginner_Vocabulary_Full.xlsx (in this folder)
  sentence_bank.xlsx  Learn_Ateso_Sentence_Bank_Processed.xlsx  (in this folder, optional)
  out.dart            ../lib/data/fake/seed_data.dart (writes straight into the
                       app - never leave a generated copy at the project root,
                       it breaks `flutter analyze`, which scans every .dart
                       file under the project)

Requires: pip install openpyxl

Re-run this whenever the spreadsheet changes. Only vocabulary, quizzes,
aptitude tests and (where real data exists) sentence exercises / one dialogue
are generated from real data. Everything else is clearly-labeled PLACEHOLDER.
"""
import random, re, sys, os
import openpyxl

VOCAB = sys.argv[1] if len(sys.argv) > 1 else "Learn_Ateso_Beginner_Vocabulary_Full.xlsx"
SENTS = sys.argv[2] if len(sys.argv) > 2 else "Learn_Ateso_Sentence_Bank_Processed.xlsx"
OUT = sys.argv[3] if len(sys.argv) > 3 else "../lib/data/fake/seed_data.dart"

MAX_PER_LESSON = 8
QUIZ_MAX = 6
APTITUDE_QUESTIONS = 5
APTITUDE_PASS = 4

PHRASE_LIKE = {"My name is...", "What is your name", "Nice to meet you", "I am from...",
               "I don't understand", "Can you help me", "How much is this"}
NUM_FIRST = ["One","Two","Three","Four","Five","Six","Seven","Eight","Nine","Ten"]

# level -> (level name, [(lesson-group title, selector)])
def sel_group(name, pred=lambda eng: True):
    return lambda g, eng: g == name and pred(eng)

LEVELS = [
    ("Level 1: Greetings and Basics", [
        ("Greetings", sel_group("Greetings", lambda e: e not in PHRASE_LIKE)),
        ("Useful Phrases", sel_group("Greetings", lambda e: e in PHRASE_LIKE)),
        ("Pronouns", sel_group("Pronouns")),
        ("Numbers", sel_group("Numbers", lambda e: e in NUM_FIRST)),
    ]),
    ("Level 2: Verbs, Questions and Time", [
        ("Bigger Numbers", sel_group("Numbers", lambda e: e not in NUM_FIRST)),
        ("Common Verbs", sel_group("Common verbs")),
        ("Question Words", sel_group("Question words")),
        ("Days and Time", sel_group("Days & time")),
    ]),
    ("Level 3: Family, Colors and Describing", [
        ("Family", sel_group("Family")),
        ("Colors", sel_group("Colors")),
        ("Describing Words", sel_group("Adjectives")),
    ]),
    ("Level 4: Body and Food", [
        ("Body Parts", sel_group("Body parts")),
        ("Food and Drink", sel_group("Food & drink")),
    ]),
    ("Level 5: Places, Objects and Weather", [
        ("Places", sel_group("Places")),
        ("Everyday Objects", sel_group("Objects")),
        ("Weather", sel_group("Weather")),
    ]),
]

def clean(a):
    return re.sub(r"[.\u2026]+$", "", a).strip()

def category(g, eng):
    if g == "Greetings": return "phrases" if eng in PHRASE_LIKE else "greetings"
    if g == "Numbers": return "numbers"
    if g in ("Pronouns", "Question words"): return "grammar"
    if g == "Days & time": return "time"
    if g == "Family": return "family"
    if g == "Body parts": return "medical"
    if g == "Food & drink": return "food"
    if g == "Places":
        return {"Hospital": "medical", "School": "work", "Home": "daily_life"}.get(eng, "travel")
    if g == "Adjectives": return "emotions" if eng in ("Happy", "Sad") else "daily_life"
    return "daily_life"

def dq(s):
    """Dart single-quoted string literal."""
    s = s.replace("\\", "\\\\").replace("'", "\\'").replace("$", "\\$").replace("\n", "\\n")
    return "'" + s + "'"

def dlist(items):
    return "[" + ", ".join(dq(i) for i in items) + "]"

def load_vocab(path):
    ws = openpyxl.load_workbook(path)["Beginner Words"]
    rows, group = [], None
    for r in ws.iter_rows(min_row=2, values_only=True):
        a, b, c = r[0], r[1], r[2]
        if a and not b:
            group = a
        elif b and c and str(c).strip():
            rows.append((group, b.strip(), clean(str(c))))
    return rows

def load_sentences(path):
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path)
    if "Beginner subset" not in wb.sheetnames:
        return []
    ws = wb["Beginner subset"]
    hd = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[hd.index("category")]:          # only sentences a human reviewed/tagged
            out.append((r[hd.index("english")], r[hd.index("ateso")]))
    return out

def tokens(s):
    return [t.lower() for t in re.sub(r"[.,?!\u2026]+", " ", s).split()]

def pick_distractors(rng, pool, correct, n):
    cands = [w for w in pool if w != correct]
    rng.shuffle(cands)
    return cands[:n]

def make_options(rng, correct, lesson_words, fallback_words):
    d = pick_distractors(rng, lesson_words, correct, 2)
    if len(d) < 2:
        d += pick_distractors(rng, [w for w in fallback_words if w not in d], correct, 2 - len(d))
    opts = [correct] + d
    rng.shuffle(opts)
    return opts, opts.index(correct)

def kind(word):
    return "phrase" if len(word.split()) > 1 else "word"

def main():
    vocab = load_vocab(VOCAB)
    sents = load_sentences(SENTS)
    all_words = [a for _, _, a in vocab]

    levels_dart, lessons_dart, content_dart, apt_dart = [], [], [], []
    level_vocab = {}            # level number -> list of (eng, ateso)
    stats = {"lessons": 0, "real_sentence": 0, "placeholder_sentence": 0}
    used_sents = set()

    for lvl_idx, (lvl_name, groups) in enumerate(LEVELS, start=1):
        lvl_id = f"level_{lvl_idx}"
        lesson_defs = []
        for title, sel in groups:
            items = [(g, e, a) for g, e, a in vocab if sel(g, e)]
            if not items:
                continue
            parts = max(1, -(-len(items) // MAX_PER_LESSON))
            size = -(-len(items) // parts)
            for p in range(parts):
                chunk = items[p * size:(p + 1) * size]
                t = title if parts == 1 else f"{title} {p + 1}"
                lesson_defs.append((t, chunk))
        level_vocab[lvl_idx] = [(e, a) for _, chunk in lesson_defs for _, e, a in chunk]

        levels_dart.append(
            f"    const Level(\n      id: {dq(lvl_id)},\n      name: {dq(lvl_name)},\n      order: {lvl_idx},\n"
            f"      puzzle: Puzzle(id: {dq(lvl_id + '-puzzle')}, totalPieces: {len(lesson_defs)}),\n    ),")

        lesson_lines = []
        for n, (title, chunk) in enumerate(lesson_defs, start=1):
            les_id = f"lesson_{lvl_idx}_{n}"
            stats["lessons"] += 1
            lesson_lines.append(
                f"      Lesson(id: {dq(les_id)}, levelId: {dq(lvl_id)}, title: {dq(title)}, order: {n}),")
            rng = random.Random(les_id)
            words = [a for _, _, a in chunk]

            vocab_lines = ",\n".join(
                f"        VocabularyItem(word: {dq(a)}, translation: {dq(e)}, audioRef: null, category: {dq(category(g, e))})"
                for g, e, a in chunk)

            quiz_lines = []
            for g, e, a in chunk[:QUIZ_MAX]:
                opts, idx = make_options(rng, a, words, all_words)
                quiz_lines.append(
                    f"        QuizQuestion(\n          prompt: {dq(f'Which {kind(a)} means \"{e}\"?')},\n"
                    f"          options: {dlist(opts)},\n          correctOptionIndex: {idx},\n        )")
            quiz = ",\n".join(quiz_lines)

            # sentence exercise: (1) real multi-word phrase from this lesson,
            # (2) reviewed corpus sentence containing a lesson word, (3) placeholder
            ex = None
            multi = [(e, a) for _, e, a in chunk if 2 <= len(a.split()) <= 6]
            if multi:
                e, a = multi[0]
                ex = (f'Arrange the words to say: "{e}"', tokens(a))
            else:
                wl = {w.lower() for w in words}
                for eng, ate in sents:
                    tk = tokens(ate)
                    if ate not in used_sents and 2 <= len(tk) <= 5 and wl & set(tk):
                        used_sents.add(ate)
                        ex = (f'Arrange the words to say: "{eng}"', tk)
                        break
            if ex:
                stats["real_sentence"] += 1
                correct = ex[1]
                scr = correct[:]
                for _ in range(10):
                    rng.shuffle(scr)
                    if scr != correct:
                        break
                sentence = (f"const SentenceExercise(\n        promptTranslation: {dq(ex[0])},\n"
                            f"        scrambledWords: {dlist(scr)},\n        correctOrder: {dlist(correct)},\n      )")
            else:
                stats["placeholder_sentence"] += 1
                sentence = ("const SentenceExercise(\n        promptTranslation: "
                            "'PLACEHOLDER: real sentence exercise needed for this lesson.',\n"
                            "        scrambledWords: ['[Word 2]', '[Word 1]'],\n"
                            "        correctOrder: ['[Word 1]', '[Word 2]'],\n      )")

            if lvl_idx == 1 and n == 1:
                dialogue = (
                    f"const DialogueModule(\n        id: {dq(les_id + '-dialogue')},\n        scenario: 'Greeting a neighbor',\n"
                    "        turns: [\n          DialogueTurn(speaker: 'npc', line: 'Yoga! Biyai bo jo?'),\n"
                    "          DialogueTurn(\n            speaker: 'user',\n            line: '',\n"
                    "            responseOptions: ['Ejokuna.', 'Mam.'],\n            correctResponseIndex: 0,\n          ),\n        ],\n      )")
            else:
                dialogue = (
                    f"const DialogueModule(\n        id: {dq(les_id + '-dialogue')},\n"
                    "        scenario: 'PLACEHOLDER dialogue: real script needed for this lesson',\n"
                    "        turns: [\n          DialogueTurn(speaker: 'npc', line: '[Placeholder line]'),\n"
                    "          DialogueTurn(\n            speaker: 'user',\n            line: '',\n"
                    "            responseOptions: ['[Placeholder A]', '[Placeholder B]'],\n            correctResponseIndex: 0,\n          ),\n        ],\n      )")

            content_dart.append(
                f"    {dq(les_id)}: LessonContent(\n      lessonId: {dq(les_id)},\n      vocabulary: const [\n{vocab_lines},\n      ],\n"
                "      grammarNote: 'PLACEHOLDER grammar note: real Ateso grammar explanation to be written with a fluent speaker.',\n"
                f"      quiz: const [\n{quiz},\n      ],\n      sentenceExercise: {sentence},\n      dialogue: {dialogue},\n    ),")

        lessons_dart.append(f"    {dq(lvl_id)}: const [\n" + "\n".join(lesson_lines) + "\n    ],")

    # aptitude tests: test for level N checks the content of level N-1
    for lvl_idx in range(2, len(LEVELS) + 1):
        rng = random.Random(f"apt{lvl_idx}")
        pool = level_vocab[lvl_idx - 1]
        picks = rng.sample(pool, min(APTITUDE_QUESTIONS, len(pool)))
        pool_words = [a for _, a in pool]
        qs = []
        for e, a in picks:
            opts, idx = make_options(rng, a, pool_words, all_words)
            qs.append(f"        QuizQuestion(prompt: {dq(f'Which {kind(a)} means \"{e}\"?')}, options: {dlist(opts)}, correctOptionIndex: {idx}),")
        apt_dart.append(
            f"    {dq(f'level_{lvl_idx}')}: const AptitudeTest(\n      levelId: {dq(f'level_{lvl_idx}')},\n      questions: [\n"
            + "\n".join(qs) + f"\n      ],\n      passThreshold: {min(APTITUDE_PASS, len(picks))},\n    ),")

    header = """import '../../models/aptitude_test.dart';
import '../../models/dialogue_module.dart';
import '../../models/lesson.dart';
import '../../models/lesson_content.dart';
import '../../models/level.dart';
import '../../models/puzzle.dart';
import '../../models/quiz_question.dart';
import '../../models/sentence_exercise.dart';
import '../../models/vocabulary_item.dart';

/// GENERATED by generate_seed_data.py - edit the spreadsheet and re-run
/// instead of editing this file by hand.
///
/// REAL: vocabulary (words + translations from a fluent Ateso speaker),
/// quizzes and aptitude tests (built from that vocabulary), and sentence
/// exercises / the one dialogue where real Ateso phrases exist.
/// PLACEHOLDER (clearly labeled in-app): grammar notes, most dialogues, and
/// sentence exercises for lessons with no real multi-word phrase yet.
/// audioRef is null everywhere until recordings exist (SRS 3.5).
class SeedData {
  SeedData._();

  static final List<Level> levels = [
"""
    body = (header + "\n".join(levels_dart) + "\n  ];\n\n  static final Map<String, List<Lesson>> lessonsByLevel = {\n"
            + "\n".join(lessons_dart) + "\n  };\n\n  static final Map<String, LessonContent> lessonContent = {\n"
            + "\n".join(content_dart) + "\n  };\n\n  static final Map<String, AptitudeTest> aptitudeTests = {\n"
            + "\n".join(apt_dart) + "\n  };\n}\n")
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(body)
    print(f"Wrote {OUT}: {len(vocab)} words, {len(LEVELS)} levels, {stats['lessons']} lessons")
    print(f"Sentence exercises from real data: {stats['real_sentence']}, placeholder: {stats['placeholder_sentence']}")

main()
